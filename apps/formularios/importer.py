"""People list upload (§7.1): CSV or Excel, email is the only required column."""

import csv
import io

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.utils.text import slugify
from django.utils.translation import gettext as _
from openpyxl import load_workbook

from apps.catalogs.models import GrupoImpactado

from .models import Carga, Persona

FIELDS = ["email", "nombre", "oficina", "sucursal", "departamento", "supervisor", "cargo"]
GROUP_SOURCES = ["departamento", "oficina", "sucursal", "cargo"]

# Accepted header spellings, compared after slugify.
ALIASES = {
    "email": {"email", "correo", "correo-electronico", "e-mail", "mail"},
    "nombre": {"nombre", "nombre-completo", "name"},
    "oficina": {"oficina"},
    "sucursal": {"sucursal"},
    "departamento": {"departamento", "depto", "area"},
    "supervisor": {"supervisor", "jefe", "lider", "correo-supervisor", "correo-del-supervisor"},
    "cargo": {"cargo", "puesto"},
}

MAX_ROWS = 20000


class PeopleImportError(Exception):
    pass


def read_rows(uploaded):
    """Returns (headers, rows) as lists of strings from a CSV or XLSX upload."""
    name = uploaded.name.lower()
    if name.endswith(".xlsx"):
        try:
            wb = load_workbook(uploaded, read_only=True, data_only=True)
        except Exception:
            raise PeopleImportError(_("No se pudo leer el archivo Excel.")) from None
        values = [
            ["" if v is None else str(v).strip() for v in row]
            for row in wb.active.iter_rows(values_only=True)
        ]
    elif name.endswith(".csv"):
        raw = uploaded.read()
        for encoding in ("utf-8-sig", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:  # e.g. a single column of emails: no delimiter to detect
            dialect = "excel"
        values = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    else:
        raise PeopleImportError(_("Sube un archivo .csv o .xlsx."))
    values = [row for row in values if any(row)]
    if not values:
        raise PeopleImportError(_("El archivo está vacío."))
    if len(values) - 1 > MAX_ROWS:
        raise PeopleImportError(_("Máximo %(n)s filas por carga.") % {"n": MAX_ROWS})
    return values[0], values[1:]


def map_headers(headers):
    """{field: column index} for the recognised headers."""
    mapping = {}
    for idx, header in enumerate(headers):
        key = slugify(header)
        for field, aliases in ALIASES.items():
            if key in aliases and field not in mapping:
                mapping[field] = idx
    if "email" not in mapping:
        raise PeopleImportError(_("Falta la columna de correo (por ejemplo «correo» o «email»)."))
    return mapping


def parse(rows, mapping):
    """Returns (records, errors). Validates emails and drops duplicates within the file."""
    records, errors, seen = [], [], set()
    for line, row in enumerate(rows, start=2):
        rec = {f: (row[i] if i < len(row) else "") for f, i in mapping.items()}
        email = rec["email"].lower()
        try:
            validate_email(email)
        except ValidationError:
            errors.append(_("Fila %(n)s: correo no válido «%(e)s».") % {"n": line, "e": email})
            continue
        if email in seen:
            errors.append(
                _("Fila %(n)s: correo repetido %(e)s (se usa la primera fila).")
                % {"n": line, "e": email}
            )
            continue
        seen.add(email)
        rec["email"] = email
        records.append(rec)
    return records, errors


@transaction.atomic
def apply(records, *, filename, user, group_field="", create_groups=False, errors=()):
    """Creates or updates people by email. Returns the Carga."""
    groups = {g.name.lower(): g for g in GrupoImpactado.objects.all()}
    created = updated = 0
    for rec in records:
        values = {f: rec.get(f, "") for f in FIELDS if f != "email" and f in rec}
        if group_field:
            name = " ".join(rec.get(group_field, "").split())
            if name:
                grupo = groups.get(name.lower())
                if grupo is None and create_groups:
                    grupo = GrupoImpactado.objects.create(name=name)
                    groups[name.lower()] = grupo
                values["grupo"] = grupo
        persona = Persona.objects.filter(email__iexact=rec["email"]).first()
        if persona is None:
            Persona.objects.create(email=rec["email"], **values)
            created += 1
        else:
            for field, value in values.items():
                setattr(persona, field, value)
            persona.is_active = True
            persona.save()
            updated += 1
    return Carga.objects.create(
        archivo=filename,
        columna_grupo=group_field,
        filas=len(records) + len(errors),
        creadas=created,
        actualizadas=updated,
        errores=list(errors),
        creada_por=user,
    )
