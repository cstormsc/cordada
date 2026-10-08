import datetime
import io
import re
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.catalogs.models import GrupoImpactado
from apps.core.models import Role
from apps.core.tenancy import tenant_context
from apps.formularios.models import (
    Cuestionario,
    Enlace,
    EstadoCuestionario,
    Persona,
    Pregunta,
    Respuesta,
)

pytestmark = pytest.mark.django_db
D = Decimal


def base(c):
    return f"/c/acme/formularios/{c.numero}/"


class TestEditor:
    def add_question(self, client, c, texto="¿Qué hace el botón?"):
        return client.post(f"{base(c)}preguntas/nueva/", {
            "texto": texto, "tipo": "simple",
            "op-TOTAL_FORMS": "5", "op-INITIAL_FORMS": "0", "op-MIN_NUM_FORMS": "0",
            "op-MAX_NUM_FORMS": "5",
            "op-0-texto": "Guarda", "op-0-correcta": "on", "op-1-texto": "Borra",
        })

    def points(self, c, acme):
        with tenant_context(acme):
            return [p.puntos for p in Pregunta.objects.filter(cuestionario=c)]

    def test_automatic_split_until_edited_by_hand(self, client, equipo, make_cuestionario,
                                                  acme):
        c = make_cuestionario()
        client.force_login(equipo)
        for _ in range(3):
            self.add_question(client, c)
        assert self.points(c, acme) == [D("33.34"), D("33.33"), D("33.33")]

        with tenant_context(acme):
            ids = list(Pregunta.objects.filter(cuestionario=c).values_list("pk", flat=True))
        client.post(f"{base(c)}puntos/", {f"puntos-{ids[0]}": "50", f"puntos-{ids[1]}": "25",
                                          f"puntos-{ids[2]}": "25,00"})
        assert self.points(c, acme) == [D("50"), D("25"), D("25")]
        self.add_question(client, c)  # manual points are kept; the new one starts at 0
        assert self.points(c, acme) == [D("50"), D("25"), D("25"), D("0")]
        body = client.get(base(c)).content.decode()
        assert "tiene 0 puntos" in body

        client.post(f"{base(c)}repartir/")
        assert self.points(c, acme) == [D("25")] * 4

    def test_deleting_redistributes_in_automatic_mode(self, client, equipo, make_cuestionario,
                                                      acme):
        c = make_cuestionario(4)
        client.force_login(equipo)
        with tenant_context(acme):
            first = Pregunta.objects.filter(cuestionario=c).first()
        client.post(f"{base(c)}preguntas/{first.pk}/borrar/")
        assert self.points(c, acme) == [D("33.34"), D("33.33"), D("33.33")]

    def test_points_warning_when_not_100(self, client, equipo, make_cuestionario, acme):
        c = make_cuestionario(2)
        client.force_login(equipo)
        with tenant_context(acme):
            ids = list(Pregunta.objects.filter(cuestionario=c).values_list("pk", flat=True))
        response = client.post(f"{base(c)}puntos/", {f"puntos-{ids[0]}": "60",
                                                     f"puntos-{ids[1]}": "30"}, follow=True)
        assert "Suman 90.00" in response.content.decode()
        response = client.post(f"{base(c)}publicar/", follow=True)
        assert "deben sumar 100" in response.content.decode()
        c.refresh_from_db()
        assert c.is_draft

    def test_publish_locks_questions_and_points(self, client, equipo, make_cuestionario, acme):
        c = make_cuestionario(2)
        client.force_login(equipo)
        client.post(f"{base(c)}publicar/")
        c.refresh_from_db()
        assert c.estado == EstadoCuestionario.PUBLICADO
        with tenant_context(acme):
            p = Pregunta.objects.filter(cuestionario=c).first()
        assert client.get(f"{base(c)}preguntas/nueva/").url == base(c)
        client.post(f"{base(c)}puntos/", {f"puntos-{p.pk}": "99"})
        client.post(f"{base(c)}repartir/")
        client.post(f"{base(c)}preguntas/{p.pk}/borrar/")
        assert self.points(c, acme) == [D("50"), D("50")]
        # Only the closing date can still change.
        client.post(f"{base(c)}editar/", {"titulo": "Otro", "nota_minima": "10", "intentos": "9",
                                         "vence": "2030-01-01"})
        c.refresh_from_db()
        assert (c.titulo, c.nota_minima, c.intentos, c.vence) == (
            "Nuevo flujo", D("80"), 1, datetime.date(2030, 1, 1))

    def test_guide_designs_but_cannot_send_links(self, client, member, acme, make_cuestionario):
        c = make_cuestionario(2)
        guia, _ = member(Role.GUIA, client=acme, email="guia@example.com")
        client.force_login(guia)
        assert client.get(base(c)).status_code == 200
        client.post(f"{base(c)}publicar/")
        assert client.get(f"{base(c)}enlaces/").status_code == 403


@pytest.fixture
def published(make_cuestionario, acme):
    c = make_cuestionario(2, intentos=2)
    with tenant_context(acme):
        c.estado = EstadoCuestionario.PUBLICADO
        c.save()
        norte = GrupoImpactado.objects.create(name="Norte")
        for i in range(6):
            Persona.objects.create(email=f"p{i}@x.com", nombre=f"Persona {i}", grupo=norte)
        Persona.objects.create(email="sur@x.com", nombre="Del sur")
    return c


def links(client, c):
    ws = load_workbook(io.BytesIO(client.get(f"{base(c)}enlaces/exportar/").content)).active
    return {row[0].value: row[3].value for row in ws.iter_rows(min_row=2)}


class TestLinksAndAnswers:
    def test_generate_by_group_once_and_export(self, client, equipo, published, acme):
        client.force_login(equipo)
        with tenant_context(acme):
            norte = GrupoImpactado.objects.get(name="Norte")
        client.post(f"{base(published)}enlaces/", {"grupos": [norte.pk]})
        client.post(f"{base(published)}enlaces/", {})  # everyone: only the missing one is added
        with tenant_context(acme):
            assert Enlace.objects.count() == 7
        exported = links(client, published)
        assert len(exported) == 7
        assert all(re.match(r"http://testserver/f/[\w:-]+/$", u) for u in exported.values())

    def answer(self, client, url, right=True):
        page = client.get(url).content.decode()
        names = sorted(set(re.findall(r'name="(p\d+)"', page)))
        data = {}
        for name in names:
            values = re.findall(rf'name="{name}" value="(\d+)"', page)
            data[name] = values[0] if right else values[1]
        return client.post(url, data)

    def test_answer_scores_and_limits_attempts(self, client, equipo, published, acme):
        client.force_login(equipo)
        client.post(f"{base(published)}enlaces/", {})
        url = links(client, published)["p0@x.com"]
        client.logout()  # respondents have no account
        page = client.get(url).content.decode()
        assert "confidenciales, no anónimas" in page
        first = self.answer(client, url, right=False).content.decode()
        assert "No aprobado" in first and "Intentar de nuevo" in first
        second = self.answer(client, url).content.decode()
        assert "100" in second and "Aprobado" in second
        assert "Ya usaste todos tus intentos" in client.get(url).content.decode()
        with tenant_context(acme):
            assert Respuesta.objects.count() == 2

    def test_unanswered_questions_are_flagged(self, client, equipo, published):
        client.force_login(equipo)
        client.post(f"{base(published)}enlaces/", {})
        url = links(client, published)["p1@x.com"]
        client.logout()
        assert "Responde todas las preguntas" in client.post(url, {}).content.decode()

    def test_closed_expired_tampered_and_inactive(self, client, equipo, published, acme):
        client.force_login(equipo)
        client.post(f"{base(published)}enlaces/", {})
        all_links = links(client, published)
        client.logout()
        assert client.get(all_links["p2@x.com"][:-3] + "xx/").status_code == 404
        with tenant_context(acme):
            Persona.objects.filter(email="p3@x.com").update(is_active=False)
        assert client.get(all_links["p3@x.com"]).status_code == 404
        Cuestionario.unscoped.filter(pk=published.pk).update(vence=datetime.date(2020, 1, 1))
        assert client.get(all_links["p2@x.com"]).status_code == 410
        Cuestionario.unscoped.filter(pk=published.pk).update(vence=None, estado="cerrado")
        assert client.get(all_links["p2@x.com"]).status_code == 410

    def test_bumping_version_revokes_old_link(self, client, equipo, published, acme):
        client.force_login(equipo)
        client.post(f"{base(published)}enlaces/", {})
        old = links(client, published)["p4@x.com"]
        Enlace.unscoped.filter(persona__email="p4@x.com").update(version=2)
        assert client.get(old).status_code == 404


class TestResults:
    def respond(self, published, acme, emails, score):
        with tenant_context(acme):
            for email in emails:
                enlace, _ = Enlace.objects.get_or_create(
                    cuestionario=published, persona=Persona.objects.get(email=email))
                Respuesta.objects.create(enlace=enlace, intento=1, elecciones={},
                                         puntos_por_pregunta={}, puntaje=D(score),
                                         aprobado=D(score) >= D("80"))

    def test_small_groups_are_hidden_and_detail_only_for_change_team(
            self, client, equipo, member, published, acme):
        self.respond(published, acme, [f"p{i}@x.com" for i in range(5)], "90")
        self.respond(published, acme, ["sur@x.com"], "10")
        client.force_login(equipo)
        body = client.get(f"{base(published)}resultados/").content.decode()
        assert "100 %" in body  # Norte: 5 respondents, all passed
        assert "menos de 5" in body  # "Sin dato" group with a single respondent is hidden
        assert "Detalle individual" in body and "sur@x.com" in body

        admin, _ = member(Role.ADMINISTRACION, client=acme, email="adm@example.com")
        client.force_login(admin)
        body = client.get(f"{base(published)}resultados/").content.decode()
        assert "Detalle individual" not in body and "sur@x.com" not in body

    def test_impact_page_shows_group_verification(self, client, equipo, make_impacto,
                                                  published, acme):
        with tenant_context(acme):
            norte = GrupoImpactado.objects.get(name="Norte")
        impacto = make_impacto(grupo=norte)
        Cuestionario.unscoped.filter(pk=published.pk).update(impacto=impacto)
        self.respond(published, acme, [f"p{i}@x.com" for i in range(6)], "85")
        client.force_login(equipo)
        body = client.get(f"/c/acme/matriz/{impacto.numero}/").content.decode()
        assert "Verificación de competencias" in body and "100 %" in body

    def test_area_leader_has_no_questionnaire_access(self, client, member, acme, published):
        leader, _ = member(Role.LIDER_AREA, client=acme, email="lider@example.com")
        client.force_login(leader)
        assert client.get("/c/acme/formularios/").status_code == 403
        assert client.get(f"{base(published)}resultados/").status_code == 403
