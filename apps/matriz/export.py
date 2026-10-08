"""Matrix rows for Excel export and stage snapshots (same columns for both)."""

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("id", "ID"),
    ("titulo", "Título"),
    ("estado", "Estado"),
    ("prioridad", "Prioridad"),
    ("depende_de", "Depende de"),
    ("prioridad_ajustada", "Prioridad ajustada"),
    ("tipo", "Tipo"),
    ("tiempo", "Tiempo"),
    ("origen", "Origen"),
    ("grupo", "Grupo impactado"),
    ("personas", "Personas impactadas"),
    ("descripcion", "Descripción"),
    ("situacion_actual", "Situación actual"),
    ("situacion_futura", "Situación futura"),
    ("magnitud", "Magnitud"),
    ("resistencia", "Resistencia"),
    ("preparacion", "Preparación"),
    ("probabilidad", "Probabilidad"),
    ("severidad", "Severidad"),
    ("riesgo", "Nivel de riesgo"),
    ("riesgo_descripcion", "Riesgo"),
    ("comunicacion", "Requiere comunicación"),
    ("movilizacion", "Requiere movilización"),
    ("formacion", "Requiere formación"),
    ("cambio_conocimiento", "Cambio en el conocimiento"),
    ("formacion_criticidad", "Criticidad formación"),
    ("formacion_momento", "Momento formación"),
    ("medida", "Medida"),
    ("responsable", "Responsable"),
    ("responsable_categoria", "Categoría responsable"),
    ("fecha_compromiso", "Fecha compromiso"),
    ("criterio_evidencia", "Evidencia de cierre"),
    ("criterio_meta", "Meta"),
    ("criterio_descripcion", "Criterio de cierre"),
    ("certificacion_lider", "Requiere certificación del líder"),
]


def _yes_no(value):
    return {True: "Sí", False: "No"}.get(value, "")


def _str(value):
    return str(value) if value is not None else ""


def impacto_rows(impactos):
    """JSON-serializable rows; dates as ISO strings."""
    rows = []
    for i in impactos:
        rows.append(
            {
                "id": i.codigo,
                "titulo": i.titulo,
                "estado": i.get_estado_display(),
                "prioridad": i.get_prioridad_display() or "",
                "depende_de": ", ".join(d.depende_de.codigo for d in i.dependencias.all()),
                "prioridad_ajustada": "Sí" if i.prioridad_ajustada else "",
                "tipo": _str(i.tipo),
                "tiempo": _str(i.tiempo),
                "origen": _str(i.origen),
                "grupo": _str(i.grupo)
                or (f"(sugerido) {i.grupo_sugerido}" if i.grupo_sugerido else ""),
                "personas": i.personas_impactadas,
                "descripcion": i.descripcion,
                "situacion_actual": i.situacion_actual,
                "situacion_futura": i.situacion_futura,
                "magnitud": i.magnitud,
                "resistencia": i.resistencia,
                "preparacion": i.preparacion,
                "probabilidad": i.probabilidad,
                "severidad": i.severidad,
                "riesgo": i.get_riesgo_nivel_display() or "",
                "riesgo_descripcion": i.riesgo_descripcion,
                "comunicacion": _yes_no(i.requiere_comunicacion),
                "movilizacion": _yes_no(i.requiere_movilizacion),
                "formacion": _yes_no(i.requiere_formacion),
                "cambio_conocimiento": _str(i.tipo_cambio_conocimiento),
                "formacion_criticidad": i.get_formacion_criticidad_display(),
                "formacion_momento": _str(i.formacion_momento),
                "medida": i.medida,
                "responsable": i.responsable.display_name if i.responsable else "",
                "responsable_categoria": _str(i.responsable_categoria),
                "fecha_compromiso": i.fecha_compromiso.isoformat() if i.fecha_compromiso else "",
                "criterio_evidencia": _str(i.criterio_evidencia),
                "criterio_meta": float(i.criterio_meta) if i.criterio_meta is not None else None,
                "criterio_descripcion": i.criterio_descripcion,
                "certificacion_lider": _yes_no(i.requiere_certificacion_lider),
            }
        )
    return rows


def workbook(rows, title):
    wb = Workbook()
    ws = wb.active
    ws.title = "Matriz"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([label for _key, label in COLUMNS])
    for cell in ws[2]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([safe_cell(row.get(key)) for key, _label in COLUMNS])
    ws.freeze_panes = "C3"
    for idx, (key, _label) in enumerate(COLUMNS, start=1):
        width = 50 if key in {"descripcion", "medida", "riesgo_descripcion"} else 18
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.auto_filter.ref = f"A2:{get_column_letter(len(COLUMNS))}{len(rows) + 2}"
    return wb


def safe_cell(value):
    # Guard against spreadsheet formula injection from free text.
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value
