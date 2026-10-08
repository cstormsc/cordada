import pytest
from django.core.exceptions import ValidationError

from apps.core.models import Role
from apps.core.tenancy import tenant_context
from apps.matriz import dependencias, workflow
from apps.matriz.models import Dependencia, Estado, Impacto

pytestmark = pytest.mark.django_db


@pytest.fixture
def abc(make_impacto):
    return make_impacto(titulo="A"), make_impacto(titulo="B"), make_impacto(titulo="C")


class TestRules:
    def test_self_duplicate_and_cycles_are_refused(self, abc, acme, equipo):
        a, b, c = abc
        with tenant_context(acme):
            dependencias.add(a, b, user=equipo)
            dependencias.add(b, c, user=equipo)
            for x, y in [(a, a), (a, b), (b, a), (c, a)]:  # self, duplicate, 2- and 3-cycle
                with pytest.raises(ValidationError):
                    dependencias.add(x, y, user=equipo)
            dependencias.add(a, c, user=equipo)  # a shortcut is not a cycle
            assert Dependencia.objects.count() == 3

    @pytest.mark.parametrize(
        ("estado", "blocks"),
        [
            (Estado.PROPUESTO, True),
            (Estado.EVALUADO, True),
            (Estado.EN_ATENCION, True),
            (Estado.REABIERTO, True),
            (Estado.ATENDIDO, False),
            (Estado.CERRADO, False),
            (Estado.DESCARTADO, False),
        ],
    )
    def test_which_prerequisite_states_block(self, abc, acme, equipo, estado, blocks):
        a, b, _ = abc
        with tenant_context(acme):
            dependencias.add(a, b, user=equipo)
            Impacto.objects.filter(pk=b.pk).update(estado=estado)
            assert dependencias.open_prerequisites(a).exists() is blocks
            assert dependencias.with_blocked(
                Impacto.objects.filter(pk=a.pk)
            ).get().n_bloqueos == int(blocks)

    def test_dependencies_never_block_the_lifecycle(self, make_impacto, acme, equipo):
        a = make_impacto(complete=True, estado=Estado.EN_ATENCION)
        b = make_impacto(titulo="Prerrequisito abierto")
        with tenant_context(acme):
            dependencias.add(a, b, user=equipo)
            workflow.transition(a, Estado.ATENDIDO, user=equipo)
            warning = dependencias.transition_warning(a)
            workflow.transition(a, Estado.CERRADO, user=equipo)
        assert a.estado == Estado.CERRADO
        assert "IMP-002 (Propuesto)" in warning

    def test_no_warning_outside_attended_or_closed(self, abc, acme, equipo):
        a, b, _ = abc
        with tenant_context(acme):
            dependencias.add(a, b, user=equipo)
            assert dependencias.transition_warning(a) is None  # a is Propuesto


class TestViews:
    def url(self, impacto, suffix=""):
        return f"/c/acme/matriz/{impacto.numero}/{suffix}"

    def test_change_team_adds_and_removes(self, client, equipo, abc, acme):
        a, b, _ = abc
        client.force_login(equipo)
        client.post(self.url(a, "dependencias/"), {"depende_de": b.pk})
        with tenant_context(acme):
            dep = Dependencia.objects.get()
            assert (dep.impacto, dep.depende_de, dep.creada_por) == (a, b, equipo)
        body = client.get(self.url(a)).content.decode()
        assert "Bloqueado" in body and "Quitar" in body
        assert "IMP-001" in client.get(self.url(b)).content.decode()  # "dependen de este"
        client.post(self.url(a, f"dependencias/{dep.pk}/quitar/"))
        with tenant_context(acme):
            assert not Dependencia.objects.exists()
            assert Dependencia.history.count() == 2  # created and removed, both recorded

    def test_cycle_error_is_shown(self, client, equipo, abc):
        a, b, _ = abc
        client.force_login(equipo)
        client.post(self.url(a, "dependencias/"), {"depende_de": b.pk})
        response = client.post(self.url(b, "dependencias/"), {"depende_de": a.pk}, follow=True)
        assert "se formaría un ciclo" in response.content.decode()

    @pytest.mark.parametrize("role", [Role.ADMINISTRACION, Role.PATROCINADOR])
    def test_other_roles_see_but_cannot_change(self, client, member, acme, abc, equipo, role):
        a, b, c = abc
        with tenant_context(acme):
            dep = dependencias.add(a, b, user=equipo)
        user, _ = member(role, client=acme, email="otro@example.com")
        client.force_login(user)
        body = client.get(self.url(a)).content.decode()
        assert "Depende de" in body and "IMP-002" in body
        assert "Quitar" not in body and "Agregar: depende de" not in body
        assert client.post(self.url(a, "dependencias/"), {"depende_de": c.pk}).status_code == 403
        assert client.post(self.url(a, f"dependencias/{dep.pk}/quitar/")).status_code == 403

    def test_cannot_link_another_clients_impact(self, client, equipo, abc, make_client):
        a, _, _ = abc
        with tenant_context(make_client("globex")):
            ajeno = Impacto.objects.create(titulo="x", descripcion="y", registrado_por=equipo)
        client.force_login(equipo)
        assert (
            client.post(self.url(a, "dependencias/"), {"depende_de": ajeno.pk}).status_code == 404
        )

    def test_state_change_warns_but_proceeds(self, client, equipo, make_impacto, acme):
        a = make_impacto(complete=True, estado=Estado.EN_ATENCION)
        b = make_impacto()
        with tenant_context(acme):
            dependencias.add(a, b, user=equipo)
        client.force_login(equipo)
        body = client.post(
            self.url(a, "estado/"), {"hacia": "atendido"}, follow=True
        ).content.decode()
        assert "Estado: Atendido" in body and "depende de impactos aún abiertos" in body

    def test_board_move_returns_warning_and_card_flag(self, client, equipo, make_impacto, acme):
        a = make_impacto(complete=True, estado=Estado.EN_ATENCION)
        b = make_impacto()
        with tenant_context(acme):
            dependencias.add(a, b, user=equipo)
        client.force_login(equipo)
        assert "Bloqueado" in client.get("/c/acme/matriz/tablero/").content.decode()
        data = client.post(self.url(a, "mover/"), {"hacia": "atendido"}).json()
        assert data["ok"] is True
        assert "IMP-002" in data["warning"]

    def test_blocked_filter_and_export(self, client, equipo, abc, acme):
        import io

        from openpyxl import load_workbook

        a, b, _ = abc
        with tenant_context(acme):
            dependencias.add(a, b, user=equipo)
        client.force_login(equipo)
        body = client.get("/c/acme/matriz/?alerta=bloqueados").content.decode()
        assert ">A<" in body and ">B<" not in body
        ws = load_workbook(io.BytesIO(client.get("/c/acme/matriz/exportar/").content)).active
        header = [c.value for c in ws[2]]
        col = header.index("Depende de")
        rows = {r[0].value: r[col].value for r in ws.iter_rows(min_row=3)}
        assert rows["IMP-001"] == "IMP-002"
