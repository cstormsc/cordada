import io
import re

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from apps.catalogs.models import GrupoImpactado, Scale
from apps.core.models import Role
from apps.core.tenancy import tenant_context
from apps.matriz.models import Adjunto, Estado, Foto, Impacto

from .conftest import item

pytestmark = pytest.mark.django_db

BASE = "/c/acme/matriz/"


@pytest.fixture
def as_role(client, member, acme):
    def login(role, email=None):
        user, _ = member(role, client=acme, email=email or f"{role}@example.com")
        client.force_login(user)
        return user

    return login


@pytest.fixture
def as_equipo(client, equipo):
    client.force_login(equipo)
    return equipo


def full_post(**overrides):
    data = {
        "titulo": "Nuevo flujo",
        "descripcion": "Cambia el proceso",
        "tipo": item("tipo_impacto", "proceso").pk,
        "magnitud": "2",
        "probabilidad": "3",
        "severidad": "2",
        "requiere_comunicacion": "true",
        "requiere_movilizacion": "false",
        "requiere_formacion": "unknown",
    }
    data.update(overrides)
    return data


class TestAccess:
    @pytest.mark.parametrize("role", [Role.ADMINISTRACION, Role.EQUIPO_CAMBIO, Role.PATROCINADOR])
    def test_viewers_see_the_matrix(self, client, as_role, role):
        as_role(role)
        assert client.get(BASE).status_code == 200

    @pytest.mark.parametrize("role", [Role.CAPITAN, Role.GUIA, Role.EQUIPO_TECNICO])
    def test_proposers_go_to_the_proposal_form(self, client, as_role, role):
        as_role(role)
        assert client.get(BASE).url == f"{BASE}nuevo/"

    def test_area_leaders_have_no_access(self, client, as_role):
        as_role(Role.LIDER_AREA)
        assert client.get(BASE).status_code == 403
        assert client.get(f"{BASE}nuevo/").status_code == 403

    def test_sponsor_cannot_create_or_change_state(self, client, as_role, make_impacto):
        impacto = make_impacto(complete=True)
        as_role(Role.PATROCINADOR)
        assert client.get(f"{BASE}nuevo/").status_code == 403
        assert client.get(f"{BASE}{impacto.numero}/editar/").status_code == 403
        client.post(f"{BASE}{impacto.numero}/estado/", {"hacia": "evaluado"})
        impacto.refresh_from_db()
        assert impacto.estado == Estado.PROPUESTO

    def test_disabled_module_is_forbidden(self, client, as_equipo, acme):
        acme.enabled_modules = ["ruta"]
        acme.save()
        assert client.get(BASE).status_code == 403


class TestCreateAndEdit:
    def test_change_team_creates_with_full_form(self, client, as_equipo, grupo, acme):
        response = client.post(f"{BASE}nuevo/", full_post(grupo=grupo.pk))
        assert response.url == f"{BASE}1/"
        with tenant_context(acme):
            impacto = Impacto.objects.get()
        assert impacto.estado == Estado.PROPUESTO
        assert impacto.registrado_por == as_equipo
        assert (
            impacto.requiere_comunicacion,
            impacto.requiere_movilizacion,
            impacto.requiere_formacion,
        ) == (True, False, None)
        assert impacto.prioridad == 3  # magnitud 2 + Medio (3×2=6) 2 = 4 -> P3

    def test_cannot_pick_another_clients_group(self, client, as_equipo, make_client, acme):
        with tenant_context(make_client("globex")):
            ajeno = GrupoImpactado.objects.create(name="Ajeno")
        response = client.post(f"{BASE}nuevo/", full_post(grupo=ajeno.pk))
        assert response.status_code == 200
        with tenant_context(acme):
            assert not Impacto.objects.exists()

    def test_pending_scales_are_left_out_and_defined_ones_carry_definitions(
        self, client, as_equipo
    ):
        level = Scale.objects.get(code="magnitud").levels.get(value=4)
        level.definition = "Transforma el trabajo"
        level.save()
        body = client.get(f"{BASE}nuevo/").content.decode()
        assert 'name="magnitud"' in body and "Transforma el trabajo" in body
        assert 'name="resistencia"' not in body
        assert "escala por definir" in body

    def test_closure_target_cannot_be_negative(self, client, as_equipo, acme):
        response = client.post(f"{BASE}nuevo/", full_post(criterio_meta="-50"))
        assert response.status_code == 200 and "mayor o igual a 0" in response.content.decode()
        client.post(f"{BASE}nuevo/", full_post(criterio_meta="0"))
        with tenant_context(acme):
            assert Impacto.objects.get().criterio_meta == 0

    def test_locked_impact_must_be_reopened_first(self, client, as_equipo, make_impacto):
        impacto = make_impacto(estado=Estado.CERRADO)
        assert client.get(f"{BASE}{impacto.numero}/editar/").url == f"{BASE}{impacto.numero}/"


class TestProposals:
    def test_captain_proposes_with_suggested_group(self, client, as_role, acme):
        captain = as_role(Role.CAPITAN)
        body = client.get(f"{BASE}nuevo/").content.decode()
        assert "Nuevo grupo" not in body and 'name="magnitud"' not in body
        client.post(
            f"{BASE}nuevo/",
            {"titulo": "Turnos", "descripcion": "Cambian turnos", "grupo_sugerido": "Call center"},
        )
        with tenant_context(acme):
            impacto = Impacto.objects.get()
        assert impacto.registrado_por == captain
        assert impacto.grupo_sugerido == "Call center"

    def test_proposers_list_their_own_proposals_but_cannot_open_them(
        self, client, as_role, member, acme
    ):
        """Proposers follow the state of what they sent; only viewers open impacts."""
        as_role(Role.GUIA)
        response = client.post(f"{BASE}nuevo/", {"titulo": "Turnos", "descripcion": "d"})
        assert response.url == f"{BASE}nuevo/"
        with tenant_context(acme):
            impacto = Impacto.objects.get()
            otro, _ = member(Role.CAPITAN, client=acme, email="otro@example.com")
            Impacto.objects.create(titulo="Propuesta ajena", descripcion="d", registrado_por=otro)
        body = client.get(f"{BASE}nuevo/").content.decode()
        assert "Mis propuestas" in body and impacto.codigo in body and "Propuesto" in body
        assert "Propuesta ajena" not in body
        assert f'href="{BASE}{impacto.numero}/"' not in body
        assert client.get(f"{BASE}{impacto.numero}/").status_code == 404
        assert client.get(f"{BASE}{impacto.numero}/editar/").status_code == 404
        adjuntar = client.post(f"{BASE}{impacto.numero}/adjuntos/", {"url": "https://x.com"})
        assert adjuntar.status_code == 404

    def test_change_team_creates_suggested_group(self, client, as_equipo, make_impacto, acme):
        impacto = make_impacto(grupo_sugerido="Call Center")
        client.post(f"{BASE}{impacto.numero}/grupo-sugerido/")
        impacto.refresh_from_db()
        assert (impacto.grupo.name, impacto.grupo_sugerido) == ("Call Center", "")

    def test_change_team_maps_suggestion_to_existing_group(
        self, client, as_equipo, make_impacto, grupo
    ):
        impacto = make_impacto(grupo_sugerido="cobranza")
        client.post(f"{BASE}{impacto.numero}/grupo-sugerido/", {"grupo": grupo.pk})
        impacto.refresh_from_db()
        assert impacto.grupo == grupo


class TestInlineGroup:
    url = f"{BASE}grupos/nuevo/"

    def test_creates_and_selects(self, client, as_equipo, acme):
        response = client.post(
            self.url, {"nuevo_grupo": "  Mesa   de ayuda "}, HTTP_HX_REQUEST="true"
        )
        with tenant_context(acme):
            grupo = GrupoImpactado.objects.get()
        assert grupo.name == "Mesa de ayuda"
        body = response.content.decode()
        assert f'value="{grupo.pk}" selected' in body
        assert "creado y seleccionado" in body

    def test_existing_name_is_reused(self, client, as_equipo, grupo, acme):
        body = client.post(self.url, {"nuevo_grupo": "COBRANZAS"}).content.decode()
        assert "Ya existía" in body and f'value="{grupo.pk}" selected' in body
        with tenant_context(acme):
            assert GrupoImpactado.objects.count() == 1

    def test_empty_name_keeps_current_selection(self, client, as_equipo, grupo):
        body = client.post(self.url, {"nuevo_grupo": " ", "grupo": grupo.pk}).content.decode()
        assert "Escribe el nombre" in body and f'value="{grupo.pk}" selected' in body

    @pytest.mark.parametrize("role", [Role.CAPITAN, Role.GUIA, Role.PATROCINADOR])
    def test_only_editors_create_groups(self, client, as_role, role, acme):
        as_role(role)
        assert client.post(self.url, {"nuevo_grupo": "X"}).status_code == 403
        with tenant_context(acme):
            assert not GrupoImpactado.objects.exists()


class TestDetailActions:
    def test_state_change_shows_missing_fields(self, client, as_equipo, make_impacto):
        impacto = make_impacto()
        response = client.post(
            f"{BASE}{impacto.numero}/estado/", {"hacia": "evaluado"}, follow=True
        )
        assert "Faltan campos" in response.content.decode()

    def test_evaluate_complete_impact(self, client, as_equipo, make_impacto):
        impacto = make_impacto(complete=True)
        client.post(f"{BASE}{impacto.numero}/estado/", {"hacia": "evaluado"})
        impacto.refresh_from_db()
        assert impacto.estado == Estado.EVALUADO

    def test_priority_override_needs_justification(self, client, as_equipo, make_impacto):
        impacto = make_impacto(complete=True)
        url = f"{BASE}{impacto.numero}/prioridad/"
        client.post(url, {"prioridad_ajustada": "1", "prioridad_justificacion": ""})
        impacto.refresh_from_db()
        assert impacto.prioridad_ajustada is None
        client.post(url, {"prioridad_ajustada": "1", "prioridad_justificacion": "Regulatorio"})
        impacto.refresh_from_db()
        assert (impacto.prioridad, impacto.prioridad_justificacion) == (1, "Regulatorio")

    def test_comment_and_history(self, client, as_equipo, make_impacto, acme):
        impacto = make_impacto()
        client.post(f"{BASE}{impacto.numero}/comentarios/", {"texto": "Revisar con TI"})
        client.post(f"{BASE}{impacto.numero}/editar/", full_post(titulo="Otro título"))
        body = client.get(f"{BASE}{impacto.numero}/").content.decode()
        assert "Revisar con TI" in body
        assert "título" in body  # changed field listed in the history

    def test_attachments_upload_and_protected_download(
        self, client, as_equipo, make_impacto, member, acme, make_client
    ):
        impacto = make_impacto()
        url = f"{BASE}{impacto.numero}/adjuntos/"
        client.post(url, {"archivo": SimpleUploadedFile("plan.pdf", b"%PDF-1.4 test")})
        client.post(url, {"url": "https://example.com/doc", "nombre": "Documento"})
        client.post(url, {"archivo": SimpleUploadedFile("x.exe", b"MZ")})  # rejected
        empty = client.post(url, {}, follow=True).content.decode()  # rejected
        assert "Sube un archivo o pega un enlace" in empty and "xor" not in empty
        with tenant_context(acme):
            archivo = Adjunto.objects.get(url="")
            assert Adjunto.objects.count() == 2
        assert archivo.nombre == "plan.pdf"
        assert f"clientes/{acme.pk}/" in archivo.archivo.name

        download = f"{url}{archivo.pk}/"
        response = client.get(download)
        assert b"".join(response.streaming_content) == b"%PDF-1.4 test"

        outsider, _ = member(
            Role.EQUIPO_CAMBIO, client=make_client("globex"), email="fuera@example.com"
        )
        client.force_login(outsider)
        assert client.get(download).status_code == 404
        client.logout()
        assert client.get(download).status_code == 302  # to login


class TestListsAndExports:
    def test_filters(self, client, as_equipo, make_impacto):
        make_impacto(titulo="Sin criterio P1", magnitud=4, probabilidad=4, severidad=4)
        make_impacto(complete=True, titulo="Completo")
        body = client.get(f"{BASE}lista/?alerta=p1_sin_criterio").content.decode()
        assert "Sin criterio P1" in body and "Completo" not in body
        body = client.get(f"{BASE}lista/?q=complet").content.decode()
        assert "Completo" in body and "Sin criterio P1" not in body

    def test_dropdowns_filter_on_change_on_list_and_board(self, client, as_equipo):
        pattern = r'name="(\w+)"[^>]*onchange="this.form.submit\(\)"'
        lista = re.findall(pattern, client.get(f"{BASE}lista/").content.decode())
        assert sorted(lista) == ["alerta", "estado", "grupo", "prioridad", "tiempo"]
        tablero = re.findall(pattern, client.get(BASE).content.decode())
        assert sorted(tablero) == [
            "alerta",
            "grupo",
            "prioridad",
            "tiempo",
        ]  # no estado column filter

    def test_excel_export(self, client, as_equipo, make_impacto):
        make_impacto(complete=True, titulo='=HYPERLINK("x")')
        response = client.get(f"{BASE}exportar/")
        ws = load_workbook(io.BytesIO(response.content)).active
        headers = [c.value for c in ws[2]]
        assert headers[:4] == ["ID", "Título", "Estado", "Prioridad"]
        row = [c.value for c in ws[3]]
        assert row[0] == "IMP-001"
        assert row[1].startswith("'=")  # formula neutralised

    def test_stage_snapshot(self, client, as_equipo, make_impacto, acme):
        make_impacto(complete=True)
        cumbre = item("tiempo", "cumbre")
        client.post(f"{BASE}fotos/", {"tiempo": cumbre.pk, "nota": "Cierre Cumbre"})
        with tenant_context(acme):
            foto = Foto.objects.get()
            # Later edits do not change the snapshot.
            Impacto.objects.update(titulo="Cambiado")
        assert foto.datos[0]["titulo"] == "Nuevo flujo de cobranza"
        ws = load_workbook(
            io.BytesIO(client.get(f"{BASE}fotos/{foto.pk}/exportar/").content)
        ).active
        assert ws.cell(3, 2).value == "Nuevo flujo de cobranza"

    def test_sponsor_comments_but_cannot_edit_or_attach(self, client, as_role, make_impacto, acme):
        impacto = make_impacto(complete=True)
        as_role(Role.PATROCINADOR)
        client.post(f"{BASE}{impacto.numero}/comentarios/", {"texto": "Revisado en comité"})
        with tenant_context(acme):
            assert impacto.comentarios.get().texto == "Revisado en comité"
        assert client.get(f"{BASE}{impacto.numero}/editar/").status_code == 403
        adjuntar = client.post(f"{BASE}{impacto.numero}/adjuntos/", {"url": "https://x.com"})
        assert adjuntar.status_code == 403
        with tenant_context(acme):
            assert not impacto.adjuntos.exists()

    def test_sponsor_cannot_take_snapshots(self, client, as_role):
        as_role(Role.PATROCINADOR)
        cumbre = item("tiempo", "cumbre")
        assert client.post(f"{BASE}fotos/", {"tiempo": cumbre.pk}).status_code == 403
        assert client.get(f"{BASE}fotos/").status_code == 200
