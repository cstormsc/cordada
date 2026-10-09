import datetime
import io

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.bitacora import metrics
from apps.bitacora.models import EstadoTicket as E
from apps.bitacora.models import Ticket
from apps.catalogs.models import GrupoImpactado, SystemCatalogItem
from apps.core.models import Membership, Role
from apps.core.tenancy import tenant_context
from apps.matriz.models import Estado, Impacto

from .conftest import categoria

pytestmark = pytest.mark.django_db
BASE = "/c/acme/bitacora/"


class TestAccess:
    @pytest.mark.parametrize("role", [Role.PATROCINADOR, Role.LIDER_AREA])
    def test_roles_without_bitacora(self, client, login_as, role):
        login_as(role)
        assert client.get(BASE).status_code == 403

    def test_captain_registers_and_sees_only_own(self, client, login_as, make_ticket, acme):
        ajeno = make_ticket(titulo="De otra persona")
        captain = login_as(Role.CAPITAN)
        with tenant_context(acme):
            grupo = GrupoImpactado.objects.create(name="Norte")
            Membership.objects.get(user=captain).grupos.set([grupo])
        body = client.get(f"{BASE}nuevo/").content.decode()
        assert f'value="{grupo.pk}" selected' in body  # their only group is preselected
        assert 'name="impacto"' not in body  # captains do not see the matrix
        client.post(
            f"{BASE}nuevo/",
            {
                "titulo": "Se cae al guardar",
                "descripcion": "x",
                "categoria": categoria("tecnico").pk,
                "urgencia": "grave",
                "grupo": grupo.pk,
            },
        )
        body = client.get(BASE).content.decode()
        assert "Se cae al guardar" in body and "De otra persona" not in body
        assert client.get(f"{BASE}{ajeno.numero}/").status_code == 404
        assert client.get(f"{BASE}metricas/").status_code == 403
        assert client.get(f"{BASE}ti/").status_code == 403

    def test_captain_derives_own_severe_failure(self, client, login_as, acme):
        captain = login_as(Role.CAPITAN)
        client.post(
            f"{BASE}nuevo/",
            {
                "titulo": "Caído",
                "descripcion": "x",
                "categoria": categoria("tecnico").pk,
                "urgencia": "grave",
            },
        )
        with tenant_context(acme):
            t = Ticket.objects.get(registrado_por=captain)
        body = client.get(f"{BASE}{t.numero}/").content.decode()
        assert "Fallo grave: derivar a TI ya" in body
        client.post(f"{BASE}{t.numero}/estado/", {"hacia": "derivado_ti"})
        t.refresh_from_db()
        assert t.estado == E.DERIVADO_TI and not t.validado


class TestClassification:
    def test_guide_classifies_and_marks_recurrence(self, client, login_as, make_ticket):
        original = make_ticket(titulo="Original")
        repetido = make_ticket(titulo="Repetido")
        login_as(Role.GUIA)
        response = client.post(
            f"{BASE}{repetido.numero}/editar/",
            {
                "titulo": "Repetido",
                "descripcion": "d",
                "categoria": categoria("conocimiento").pk,
                "urgencia": "normal",
                "recurrente_de": original.pk,
            },
        )
        assert response.status_code == 302
        repetido.refresh_from_db()
        assert repetido.recurrente_de == original

    def test_technical_team_sets_client_case_number(self, client, login_as, make_ticket):
        t = make_ticket("tecnico", estado=E.DERIVADO_TI)
        login_as(Role.EQUIPO_TECNICO)
        client.post(f"{BASE}{t.numero}/caso/", {"caso_cliente": "INC-4471"})
        t.refresh_from_db()
        assert t.caso_cliente == "INC-4471"
        assert client.get(f"{BASE}{t.numero}/editar/").status_code == 403

    def test_it_live_view_lists_derived_severe_first(self, client, login_as, make_ticket):
        make_ticket("tecnico", titulo="Normal antiguo", estado=E.DERIVADO_TI)
        make_ticket("tecnico", titulo="Grave reciente", estado=E.DERIVADO_TI, urgencia="grave")
        make_ticket("tecnico", titulo="Aún nuevo")
        login_as(Role.EQUIPO_TECNICO)
        body = client.get(f"{BASE}ti/").content.decode()
        assert body.index("Grave reciente") < body.index("Normal antiguo")
        assert "Aún nuevo" not in body
        assert 'hx-trigger="every 15s"' in body
        partial = client.get(f"{BASE}ti/", HTTP_HX_REQUEST="true").content.decode()
        assert "<html" not in partial and "Grave reciente" in partial


class TestProposeImpact:
    def test_unanticipated_impact_becomes_a_proposal(
        self, client, login_as, make_ticket, acme, grupo
    ):
        SystemCatalogItem.objects.create(catalog="origen", code="bitacora", label="Bitácora")
        t = make_ticket("impacto_no_anticipado", titulo="Nadie aprueba facturas", grupo=grupo)
        user = login_as(Role.GUIA)
        response = client.post(f"{BASE}{t.numero}/proponer-impacto/")
        with tenant_context(acme):
            impacto = Impacto.objects.get()
        assert response.url == f"{BASE}{t.numero}/"  # guides don't see the matrix
        assert (impacto.estado, impacto.grupo, impacto.origen.code, impacto.registrado_por) == (
            Estado.PROPUESTO,
            grupo,
            "bitacora",
            user,
        )
        assert t.codigo in impacto.descripcion
        t.refresh_from_db()
        assert t.impacto == impacto
        assert client.post(f"{BASE}{t.numero}/proponer-impacto/").status_code == 403  # once

    def test_only_for_unanticipated_impact_tickets(self, client, login_as, make_ticket):
        t = make_ticket("tecnico")
        login_as(Role.GUIA)
        assert client.post(f"{BASE}{t.numero}/proponer-impacto/").status_code == 403


class TestMetrics:
    def test_numbers(self, make_ticket, acme, equipo):
        with tenant_context(acme):
            norte = GrupoImpactado.objects.create(name="Norte", personas=50)
        base = timezone.now() - datetime.timedelta(hours=10)
        a = make_ticket("tecnico", grupo=norte)
        b = make_ticket("tecnico", grupo=norte)
        make_ticket("conocimiento", recurrente_de=a)
        with tenant_context(acme):
            Ticket.objects.filter(pk__in=[a.pk, b.pk]).update(
                created_at=base,
                primera_respuesta_at=base + datetime.timedelta(hours=2),
                resuelto_at=base + datetime.timedelta(hours=6),
                estado=E.RESUELTO,
            )
            m = metrics.metricas()
        tec = next(r for r in m["categorias"] if r["categoria"].code == "tecnico")
        assert (tec["total"], tec["abiertos"], tec["primera_h"], tec["resolucion_h"]) == (
            2,
            0,
            2.0,
            6.0,
        )
        row = next(r for r in m["grupos"] if r["grupo"] == norte)
        assert (row["total"], row["por_100"]) == (2, 4.0)
        assert (m["recurrentes"], m["pct_recurrentes"]) == (1, 33)
        assert list(m["mas_repetidos"]) == [a]

    def test_page_and_export(self, client, login_as, make_ticket):
        make_ticket("tecnico", titulo="=cmd")
        login_as(Role.GUIA)
        assert client.get(f"{BASE}metricas/?desde=2026-01-01").status_code == 200
        import io

        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(client.get(f"{BASE}exportar/").content)).active
        assert ws.cell(2, 1).value == "TCK-001" and ws.cell(2, 2).value == "'=cmd"


class TestCaptainsAndGroups:
    """§2 confirmed: captains see their groups' tickets and say 'a mí también'."""

    def captain_of(self, login_as, grupo):
        user = login_as(Role.CAPITAN)
        Membership.objects.get(user=user).grupos.add(grupo)
        return user

    def test_sees_own_and_group_tickets_only(self, client, login_as, make_ticket, grupo, acme):
        de_su_grupo = make_ticket(titulo="No carga el reporte", grupo=grupo)
        with tenant_context(acme):
            otro = GrupoImpactado.objects.create(name="Otro")
        ajeno = make_ticket(titulo="Problema ajeno", grupo=otro)
        self.captain_of(login_as, grupo)
        body = client.get(BASE).content.decode()
        assert "No carga el reporte" in body and "Problema ajeno" not in body
        assert client.get(f"{BASE}{de_su_grupo.numero}/").status_code == 200
        assert client.get(f"{BASE}{ajeno.numero}/").status_code == 404
        # Seeing it is not editing it.
        assert client.get(f"{BASE}{de_su_grupo.numero}/editar/").status_code == 403

    def test_me_too_with_details_and_undo(self, client, login_as, make_ticket, grupo, acme):
        t = make_ticket(titulo="No carga el reporte", grupo=grupo)
        captain = self.captain_of(login_as, grupo)
        url = f"{BASE}{t.numero}/tambien/"
        client.post(url, {"texto": "A mí me pasa en la sucursal sur"})
        client.post(url, {})  # twice counts once
        with tenant_context(acme):
            assert list(t.afectados.values_list("usuario", flat=True)) == [captain.pk]
            assert t.comentarios.get().texto == "A mí me pasa en la sucursal sur"
        assert "+1" in client.get(BASE).content.decode()
        assert "1 persona más reporta" in client.get(f"{BASE}{t.numero}/").content.decode()
        client.post(url, {"quitar": "1"})
        with tenant_context(acme):
            assert not t.afectados.exists()

    def test_not_on_own_or_closed_tickets(self, client, login_as, make_ticket, grupo, acme):
        captain = self.captain_of(login_as, grupo)
        propio = make_ticket(grupo=grupo, registrado_por=captain)
        cerrado = make_ticket(grupo=grupo, estado=E.CERRADO)
        assert client.post(f"{BASE}{propio.numero}/tambien/").status_code == 403
        assert client.post(f"{BASE}{cerrado.numero}/tambien/").status_code == 403

    def test_export_counts_who_else_has_it(self, client, login_as, make_ticket, grupo, acme):
        t = make_ticket(grupo=grupo)
        self.captain_of(login_as, grupo)
        client.post(f"{BASE}{t.numero}/tambien/")
        login_as(Role.EQUIPO_CAMBIO, email="otro-equipo@example.com")
        ws = load_workbook(io.BytesIO(client.get(f"{BASE}exportar/").content)).active
        headers = [c.value for c in ws[1]]
        assert ws.cell(2, headers.index("También les pasa") + 1).value == 1
