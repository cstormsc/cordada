import datetime
import io
import re

import pytest
from openpyxl import load_workbook

from apps.catalogs.models import GrupoImpactado, SystemCatalogItem
from apps.core.models import Membership, Role
from apps.core.tenancy import tenant_context
from apps.matriz.models import Estado
from apps.ruta import reports
from apps.ruta.models import Actividad, Cumplimiento, EstadoActividad, cumplimiento

pytestmark = pytest.mark.django_db

BASE = "/c/acme/ruta/"
TODAY = datetime.date(2026, 10, 15)
D = datetime.date


def frente(code):
    return SystemCatalogItem.objects.get(catalog="frente", code=code)


@pytest.fixture
def make_actividad(acme, equipo):
    def make(grupos=(), impactos=(), **kwargs):
        data = {
            "titulo": "Reunión de equipo",
            "frente": frente("comunicacion"),
            "fecha_planificada": D(2026, 10, 1),
            "registrado_por": equipo,
        }
        data.update(kwargs)
        with tenant_context(acme):
            actividad = Actividad.objects.create(**data)
            actividad.grupos.set(grupos)
            actividad.impactos.set(impactos)
        return actividad

    return make


class TestCumplimiento:
    CASES = [
        # no_planificada, estado, planificada, real -> expected
        (True, "completado", None, D(2026, 10, 3), Cumplimiento.NO_PLANIFICADA),
        (False, "completado", D(2026, 10, 5), D(2026, 10, 5), Cumplimiento.A_TIEMPO),
        (False, "verificado", D(2026, 10, 5), D(2026, 10, 2), Cumplimiento.A_TIEMPO),
        (False, "completado", D(2026, 10, 5), D(2026, 10, 9), Cumplimiento.CON_RETRASO),
        (False, "en_curso", D(2026, 10, 5), None, Cumplimiento.NO_REALIZADA),
        (False, "no_iniciado", D(2026, 10, 14), None, Cumplimiento.NO_REALIZADA),
        (False, "no_iniciado", D(2026, 10, 15), None, Cumplimiento.PENDIENTE),
        (False, "en_curso", D(2026, 11, 1), None, Cumplimiento.PENDIENTE),
    ]

    @pytest.mark.parametrize(("no_plan", "estado", "plan", "real", "expected"), CASES)
    def test_python_and_database_agree(
        self, make_actividad, acme, no_plan, estado, plan, real, expected
    ):
        assert (
            cumplimiento(
                no_planificada=no_plan,
                estado=estado,
                fecha_planificada=plan,
                fecha_real=real,
                today=TODAY,
            )
            == expected
        )
        a = make_actividad(
            no_planificada=no_plan, estado=estado, fecha_planificada=plan, fecha_real=real
        )
        with tenant_context(acme):
            annotated = Actividad.objects.with_cumplimiento(TODAY).get(pk=a.pk)
        assert annotated.cumplimiento == expected


class TestCoherence:
    def test_orphans_and_uncovered_impacts(self, make_actividad, make_impacto, acme):
        cubierto = make_impacto(estado=Estado.EN_ATENCION)
        descubierto = make_impacto(estado=Estado.EVALUADO)
        make_impacto(estado=Estado.PROPUESTO)  # not active yet: not "descubierto"
        make_actividad(impactos=[cubierto])
        huerfana = make_actividad()
        with tenant_context(acme):
            assert list(Actividad.objects.huerfanas()) == [huerfana]
            assert list(reports.descubiertos()) == [descubierto]

    def test_numbers_per_client(self, make_actividad):
        assert [make_actividad().codigo, make_actividad().codigo] == ["ACT-001", "ACT-002"]


class TestReport:
    def test_counts_and_percentages(self, make_actividad, acme):
        make_actividad(estado="completado", fecha_real=D(2026, 10, 1))  # a tiempo
        make_actividad(estado="completado", fecha_real=D(2026, 10, 4))  # retraso
        make_actividad(estado="en_curso")  # no realizada
        make_actividad(fecha_planificada=D(2026, 11, 2))  # pendiente
        make_actividad(
            no_planificada=True,
            fecha_planificada=None,
            estado="completado",
            fecha_real=D(2026, 10, 6),
            frente=frente("movilizacion"),
        )
        with tenant_context(acme):
            data = reports.informe(Actividad.objects.all(), today=TODAY)
        com = next(r for r in data["rows"] if r["frente"].code == "comunicacion")
        assert (
            com["a_tiempo"],
            com["con_retraso"],
            com["no_realizada"],
            com["pendiente"],
            com["planificadas"],
        ) == (1, 1, 1, 1, 4)
        assert (com["pct_a_tiempo"], com["pct_realizadas"]) == (33, 67)
        assert data["total"]["no_planificada"] == 1
        assert data["total"]["pct_a_tiempo"] == 33

    def test_month_filter_uses_planned_or_real_date(self, make_actividad, acme):
        make_actividad(fecha_planificada=D(2026, 9, 30))
        make_actividad(fecha_planificada=D(2026, 10, 31))
        make_actividad(
            no_planificada=True,
            fecha_planificada=None,
            estado="completado",
            fecha_real=D(2026, 10, 2),
        )
        with tenant_context(acme):
            data = reports.informe(Actividad.objects.all(), mes=D(2026, 10, 1), today=TODAY)
        assert data["total"]["planificadas"] == 1
        assert data["total"]["no_planificada"] == 1


@pytest.fixture
def as_role(client, member, acme):
    def login(role, email=None):
        user, _ = member(role, client=acme, email=email or f"{role}@example.com")
        client.force_login(user)
        return user

    return login


class TestViews:
    def test_create_from_impact_prefills_links(self, client, equipo, make_impacto, acme):
        client.force_login(equipo)
        impacto = make_impacto(
            complete=True, tiempo=SystemCatalogItem.objects.get(catalog="tiempo", code="travesia")
        )
        body = client.get(f"{BASE}nueva/?impacto={impacto.numero}").content.decode()
        assert f'value="{impacto.pk}" selected' in body
        response = client.post(
            f"{BASE}nueva/",
            {
                "titulo": "Taller",
                "frente": frente("formacion").pk,
                "impactos": [impacto.pk],
                "grupos": [impacto.grupo_id],
                "fecha_planificada": "2026-11-10",
                "estado": "no_iniciado",
                "vocero": "Gerente",
            },
        )
        assert response.url == f"{BASE}1/"
        with tenant_context(acme):
            a = Actividad.objects.get()
            assert list(a.impactos.all()) == [impacto]
        assert a.vocero == ""  # only kept for communications
        assert "ACT-001" in client.get(f"/c/acme/matriz/{impacto.numero}/").content.decode()

    @pytest.mark.parametrize(
        ("data", "error"),
        [
            ({"estado": "no_iniciado"}, "Indica la fecha planificada"),
            ({"fecha_planificada": "2026-10-01", "estado": "completado"}, "Indica cuándo se hizo"),
            (
                {
                    "fecha_planificada": "2026-10-01",
                    "estado": "completado",
                    "fecha_real": "2099-01-01",
                },
                "no puede ser futura",
            ),
        ],
    )
    def test_validation(self, client, equipo, data, error):
        client.force_login(equipo)
        response = client.post(
            f"{BASE}nueva/", {"titulo": "X", "frente": frente("impactos").pk, **data}
        )
        assert error in response.content.decode()

    def test_captain_sees_only_their_groups(self, client, as_role, make_actividad, acme):
        with tenant_context(acme):
            mio = GrupoImpactado.objects.create(name="Sucursal Norte")
            otro = GrupoImpactado.objects.create(name="Sucursal Sur")
        visible = make_actividad(titulo="Para el norte", grupos=[mio])
        oculta = make_actividad(titulo="Para el sur", grupos=[otro])
        captain = as_role(Role.CAPITAN)
        with tenant_context(acme):
            Membership.objects.get(user=captain).grupos.set([mio])
        body = client.get(BASE).content.decode()
        assert "Para el norte" in body and "Para el sur" not in body
        assert client.get(f"{BASE}{visible.numero}/").status_code == 200
        assert client.get(f"{BASE}{oculta.numero}/").status_code == 404
        assert client.get(f"{BASE}nueva/").status_code == 403
        assert client.get(f"{BASE}informe/").status_code == 403

    @pytest.mark.parametrize("role", [Role.EQUIPO_TECNICO, Role.LIDER_AREA])
    def test_roles_without_ruta(self, client, as_role, role):
        as_role(role)
        assert client.get(BASE).status_code == 403

    def test_sponsor_reads_and_downloads_report(self, client, as_role, make_actividad):
        make_actividad(estado="completado", fecha_real=D(2026, 10, 1))
        as_role(Role.PATROCINADOR)
        assert client.get(f"{BASE}informe/?mes=2026-10").status_code == 200
        pdf = client.get(f"{BASE}informe/pdf/?mes=2026-10")
        assert pdf["Content-Type"] == "application/pdf"
        assert pdf.content.startswith(b"%PDF")
        assert client.get(f"{BASE}nueva/").status_code == 403

    def test_excel_export(self, client, equipo, make_actividad):
        make_actividad(
            titulo="@SUM(1)", estado=EstadoActividad.COMPLETADO, fecha_real=D(2026, 10, 1)
        )
        client.force_login(equipo)
        ws = load_workbook(io.BytesIO(client.get(f"{BASE}exportar/").content)).active
        assert ws.cell(2, 1).value == "ACT-001"
        assert ws.cell(2, 2).value == "'@SUM(1)"
        assert ws.cell(2, 11).value == "Hecha a tiempo"

    def test_filters_apply_on_change_except_the_text_search(self, client, equipo):
        client.force_login(equipo)
        body = client.get("/c/acme/ruta/").content.decode()
        auto = re.findall(r'name="(\w+)"[^>]*onchange="this.form.submit\(\)"', body)
        assert sorted(auto) == ["cumplimiento", "estado", "frente", "grupo", "huerfanas", "tiempo"]
        assert 'name="q"' in body

    def test_orphan_filter(self, client, equipo, make_actividad, make_impacto):
        make_actividad(titulo="Con impacto", impactos=[make_impacto()])
        make_actividad(titulo="Sin impacto")
        client.force_login(equipo)
        body = client.get(f"{BASE}?huerfanas=on").content.decode()
        assert "Sin impacto" in body and "Con impacto" not in body

    def test_matrix_uncovered_filter(self, client, equipo, make_impacto):
        make_impacto(titulo="Activo sin plan", estado=Estado.EVALUADO)
        make_impacto(titulo="Recién propuesto")
        client.force_login(equipo)
        body = client.get("/c/acme/matriz/?alerta=descubiertos").content.decode()
        assert "Activo sin plan" in body and "Recién propuesto" not in body


class TestPersonas:
    def test_admin_assigns_groups_to_members(self, client, as_role, member, acme):
        with tenant_context(acme):
            norte = GrupoImpactado.objects.create(name="Norte")
        captain, _ = member(Role.CAPITAN, client=acme, email="cap@example.com")
        as_role(Role.ADMINISTRACION)
        m = Membership.objects.get(user=captain)
        client.post(f"/c/acme/configuracion/personas/{m.pk}/", {f"m{m.pk}-grupos": [norte.pk]})
        with tenant_context(acme):
            assert list(m.grupos.all()) == [norte]

    def test_cannot_assign_another_clients_group(self, client, as_role, member, acme, make_client):
        with tenant_context(make_client("globex")):
            ajeno = GrupoImpactado.objects.create(name="Ajeno")
        captain, _ = member(Role.CAPITAN, client=acme, email="cap@example.com")
        as_role(Role.ADMINISTRACION)
        m = Membership.objects.get(user=captain)
        client.post(f"/c/acme/configuracion/personas/{m.pk}/", {f"m{m.pk}-grupos": [ajeno.pk]})
        with tenant_context(acme):
            assert not m.grupos.exists()
